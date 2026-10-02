"""
AURA Advanced Tools — v2
Covers:
  Email       : Send & Read Gmail via SMTP/IMAP + App Password
  Calendar    : Create meetings/reminders via .ics files (opens in Calendar/Outlook)
  File Ops    : Copy, Move, Rename files
  Multi-Launch: Open multiple apps at once
  Brightness  : Set/get screen brightness
  Wi-Fi       : Check status, connect, disconnect, list networks
  Documents   : Create Word (.docx) and Excel (.xlsx) files
  News        : Latest headlines via Google News RSS (no key needed)
  Briefing    : Full morning briefing — weather + news + battery
  Smart Home  : TP-Link Kasa (lights/plugs) + Philips Hue (basic)
"""

import os, sys, smtplib, imaplib, email, json, time, shutil, webbrowser
import urllib.parse, urllib.request, subprocess, threading
from pathlib import Path
from typing import Dict, Any, Optional, List
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email import policy
from email.parser import BytesParser
import datetime

sys.path.append(str(Path(__file__).resolve().parent.parent))
from core.safety_guard import log_action

# ── Load credentials from config ─────────────────────────────────────────────
try:
    from config import GMAIL_ADDRESS, GMAIL_APP_PASSWORD, KASA_DEVICE_IP, HUE_BRIDGE_IP, HUE_USERNAME
except ImportError:
    GMAIL_ADDRESS = GMAIL_APP_PASSWORD = KASA_DEVICE_IP = HUE_BRIDGE_IP = HUE_USERNAME = ""


# ══════════════════════════════════════════════════════════════════════════════
#  1.  EMAIL  —  Send & Read via Gmail SMTP / IMAP
# ══════════════════════════════════════════════════════════════════════════════

def _email_configured() -> bool:
    valid_addr = GMAIL_ADDRESS and "@" in GMAIL_ADDRESS and "your_email" not in GMAIL_ADDRESS
    valid_pwd  = GMAIL_APP_PASSWORD and "your_16char" not in GMAIL_APP_PASSWORD and len(GMAIL_APP_PASSWORD) >= 10
    return valid_addr and valid_pwd

def tool_send_email(to: str, subject: str, body: str) -> Dict[str, Any]:
    """
    Sends an email via Gmail SMTP using App Password.
    Requires GMAIL_ADDRESS and GMAIL_APP_PASSWORD set in .env
    """
    if not _email_configured():
        return {
            "status": "not_configured",
            "message": (
                "Gmail is not set up yet. Please open the .env file and fill in "
                "GMAIL_ADDRESS and GMAIL_APP_PASSWORD. "
                "Get your App Password at: https://myaccount.google.com/apppasswords"
            )
        }
    try:
        msg = MIMEMultipart("alternative")
        msg["From"]    = GMAIL_ADDRESS
        msg["To"]      = to.strip()
        msg["Subject"] = subject.strip()
        msg.attach(MIMEText(body, "plain", "utf-8"))

        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=15) as server:
            server.login(GMAIL_ADDRESS, GMAIL_APP_PASSWORD)
            server.sendmail(GMAIL_ADDRESS, to.strip(), msg.as_string())

        log_action("SEND_EMAIL", {"to": to, "subject": subject}, "SUCCESS")
        return {
            "status": "success",
            "message": f"Email sent to {to} with subject '{subject}'."
        }
    except smtplib.SMTPAuthenticationError:
        return {
            "status": "auth_error",
            "message": "Gmail authentication failed. Check your App Password in .env. Make sure 2FA is enabled and App Password is correct."
        }
    except Exception as e:
        return {"status": "error", "message": f"Failed to send email: {str(e)}"}


def tool_read_emails(count: int = 5, folder: str = "INBOX", unread_only: bool = True) -> Dict[str, Any]:
    """
    Reads the latest emails from Gmail inbox via IMAP.
    Returns sender, subject, date, and a short preview of each email.
    """
    if not _email_configured():
        return {
            "status": "not_configured",
            "message": (
                "Gmail not set up yet. Please fill GMAIL_ADDRESS and GMAIL_APP_PASSWORD in .env. "
                "Get App Password at: https://myaccount.google.com/apppasswords"
            )
        }
    try:
        mail = imaplib.IMAP4_SSL("imap.gmail.com", 993)
        mail.login(GMAIL_ADDRESS, GMAIL_APP_PASSWORD)
        mail.select(folder)

        search_criteria = "(UNSEEN)" if unread_only else "ALL"
        status, data = mail.search(None, search_criteria)
        mail_ids = data[0].split()

        if not mail_ids:
            mail.logout()
            return {
                "status": "success",
                "emails": [],
                "message": "No unread emails in your inbox. Inbox is clear!" if unread_only else "No emails found."
            }

        # Take the latest `count` emails
        recent_ids = mail_ids[-count:][::-1]
        emails = []
        for mid in recent_ids:
            _, msg_data = mail.fetch(mid, "(RFC822)")
            raw = msg_data[0][1]
            msg = BytesParser(policy=policy.default).parsebytes(raw)

            sender  = str(msg["from"])[:80]
            subject = str(msg["subject"] or "(No Subject)")[:100]
            date    = str(msg["date"] or "")[:50]

            # Get body preview
            body_preview = ""
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        try:
                            body_preview = part.get_content()[:200].strip()
                        except Exception:
                            pass
                        break
            else:
                try:
                    body_preview = msg.get_content()[:200].strip()
                except Exception:
                    pass

            emails.append({
                "from": sender,
                "subject": subject,
                "date": date,
                "preview": body_preview
            })

        mail.logout()
        log_action("READ_EMAILS", {"count": len(emails)}, "SUCCESS")

        summary_parts = [f"{i+1}. From {e['from']}: \"{e['subject']}\"" for i, e in enumerate(emails)]
        summary = f"You have {len(emails)} {'unread' if unread_only else ''} emails: " + "; ".join(summary_parts)

        return {
            "status": "success",
            "count": len(emails),
            "emails": emails,
            "message": summary
        }
    except imaplib.IMAP4.error as e:
        return {"status": "auth_error", "message": f"Gmail IMAP login failed: {str(e)}. Check your App Password."}
    except Exception as e:
        return {"status": "error", "message": f"Failed to read emails: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
#  2.  CALENDAR  —  Create meetings/reminders via .ics (opens in Windows Calendar/Outlook)
# ══════════════════════════════════════════════════════════════════════════════

def tool_create_calendar_event(
    title: str,
    date_str: str,
    time_str: str = "09:00",
    duration_minutes: int = 60,
    description: str = "",
    location: str = ""
) -> Dict[str, Any]:
    """
    Creates a calendar event as a .ics file and opens it in Windows Calendar / Outlook.
    date_str: 'YYYY-MM-DD' or 'today', 'tomorrow', 'Monday', etc.
    time_str: 'HH:MM' in 24h format
    """
    try:
        # Parse date
        today = datetime.date.today()
        date_lower = date_str.lower().strip()

        weekdays = ["monday","tuesday","wednesday","thursday","friday","saturday","sunday"]
        if date_lower == "today":
            event_date = today
        elif date_lower == "tomorrow":
            event_date = today + datetime.timedelta(days=1)
        elif date_lower in weekdays:
            target_wd = weekdays.index(date_lower)
            days_ahead = (target_wd - today.weekday()) % 7
            if days_ahead == 0:
                days_ahead = 7
            event_date = today + datetime.timedelta(days=days_ahead)
        else:
            # Try parsing YYYY-MM-DD
            event_date = datetime.datetime.strptime(date_str.strip(), "%Y-%m-%d").date()

        # Parse time
        h, m = map(int, time_str.strip().replace(".", ":").split(":"))
        start_dt = datetime.datetime(event_date.year, event_date.month, event_date.day, h, m)
        end_dt   = start_dt + datetime.timedelta(minutes=duration_minutes)

        # Format for .ics (UTC-like, local time with no TZ = floating)
        fmt = "%Y%m%dT%H%M%S"
        uid = f"aura-{int(time.time())}@aura-assistant"

        ics_content = "\r\n".join([
            "BEGIN:VCALENDAR",
            "VERSION:2.0",
            "PRODID:-//AURA Voice OS//EN",
            "CALSCALE:GREGORIAN",
            "METHOD:PUBLISH",
            "BEGIN:VEVENT",
            f"UID:{uid}",
            f"DTSTART:{start_dt.strftime(fmt)}",
            f"DTEND:{end_dt.strftime(fmt)}",
            f"SUMMARY:{title}",
            f"DESCRIPTION:{description}",
            f"LOCATION:{location}",
            f"DTSTAMP:{datetime.datetime.now().strftime(fmt)}",
            "BEGIN:VALARM",
            "TRIGGER:-PT15M",
            "ACTION:DISPLAY",
            f"DESCRIPTION:Reminder: {title}",
            "END:VALARM",
            "END:VEVENT",
            "END:VCALENDAR"
        ]) + "\r\n"

        # Save .ics to Desktop/AURA_Calendar/
        cal_dir = Path.home() / "Desktop" / "AURA_Calendar"
        cal_dir.mkdir(exist_ok=True)
        safe_title = "".join(c for c in title if c.isalnum() or c in " _-").strip()[:40]
        ics_path = cal_dir / f"{safe_title}_{start_dt.strftime('%Y%m%d_%H%M')}.ics"
        ics_path.write_text(ics_content, encoding="utf-8")

        # Open with default calendar app (Windows Calendar, Outlook, etc.)
        os.startfile(str(ics_path))
        log_action("CALENDAR_EVENT", {"title": title, "date": str(event_date), "time": time_str}, "SUCCESS")

        return {
            "status": "success",
            "event": {
                "title": title,
                "date": str(event_date),
                "time": time_str,
                "duration_minutes": duration_minutes
            },
            "message": (
                f"Meeting '{title}' scheduled for {event_date.strftime('%d %B %Y')} at {time_str} "
                f"({duration_minutes} mins). Calendar app is opening to confirm."
            )
        }
    except ValueError as e:
        return {"status": "error", "message": f"Date/time format error: {str(e)}. Use 'YYYY-MM-DD' for date and 'HH:MM' for time."}
    except Exception as e:
        return {"status": "error", "message": f"Failed to create calendar event: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
#  3.  FILE OPERATIONS  —  Copy, Move, Rename
# ══════════════════════════════════════════════════════════════════════════════

def tool_file_operation(
    action: str,
    source_path: str,
    destination: str = "",
    new_name: str = ""
) -> Dict[str, Any]:
    """
    Performs file operations: copy, move, or rename.
    action: 'copy', 'move', 'rename'
    source_path: full path of the source file
    destination: target folder path (for copy/move)
    new_name: new filename including extension (for rename)
    """
    src = Path(source_path.strip())
    action = action.lower().strip()

    if not src.exists():
        return {"status": "error", "message": f"File not found: {source_path}. Use 'find files' to locate it first."}

    try:
        if action == "copy":
            if not destination:
                return {"status": "error", "message": "Please specify a destination folder for copy."}
            dest_dir = Path(destination.strip())
            dest_dir.mkdir(parents=True, exist_ok=True)
            dest_path = dest_dir / src.name
            shutil.copy2(str(src), str(dest_path))
            log_action("FILE_COPY", {"src": str(src), "dest": str(dest_path)}, "SUCCESS")
            return {"status": "success", "message": f"Copied '{src.name}' to '{dest_dir}'."}

        elif action == "move":
            if not destination:
                return {"status": "error", "message": "Please specify a destination folder for move."}
            dest_dir = Path(destination.strip())
            dest_dir.mkdir(parents=True, exist_ok=True)
            dest_path = dest_dir / src.name
            shutil.move(str(src), str(dest_path))
            log_action("FILE_MOVE", {"src": str(src), "dest": str(dest_path)}, "SUCCESS")
            return {"status": "success", "message": f"Moved '{src.name}' to '{dest_dir}'."}

        elif action == "rename":
            if not new_name:
                return {"status": "error", "message": "Please specify a new name for the file (including extension)."}
            new_path = src.parent / new_name.strip()
            if new_path.exists():
                return {"status": "error", "message": f"A file named '{new_name}' already exists in the same folder."}
            src.rename(new_path)
            log_action("FILE_RENAME", {"src": str(src), "new": str(new_path)}, "SUCCESS")
            return {"status": "success", "message": f"Renamed '{src.name}' to '{new_name}'."}

        else:
            return {"status": "error", "message": f"Unknown action '{action}'. Use: copy, move, rename"}

    except PermissionError:
        return {"status": "error", "message": f"Permission denied. Cannot {action} '{src.name}'. It may be open in another app."}
    except Exception as e:
        return {"status": "error", "message": f"File {action} failed: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
#  4.  MULTI-APP LAUNCHER  —  Open multiple apps simultaneously
# ══════════════════════════════════════════════════════════════════════════════

def tool_open_multiple_apps(app_list: List[str]) -> Dict[str, Any]:
    """
    Opens multiple apps at once.
    app_list: list of app names e.g. ['chrome', 'spotify', 'whatsapp']
    """
    from core.os_tools import tool_open_file_or_app

    if not app_list:
        return {"status": "error", "message": "No apps specified."}

    results = []
    for app in app_list:
        result = tool_open_file_or_app(app.strip())
        results.append({"app": app, "status": result.get("status"), "msg": result.get("message", "")})
        time.sleep(0.5)  # Small delay between launches

    success_count = sum(1 for r in results if r["status"] == "success")
    app_names = ", ".join(app_list)
    log_action("MULTI_LAUNCH", {"apps": app_list, "success": success_count}, "SUCCESS")
    return {
        "status": "success",
        "results": results,
        "message": f"Launched {success_count}/{len(app_list)} apps: {app_names}."
    }


# ══════════════════════════════════════════════════════════════════════════════
#  5.  BRIGHTNESS  —  Set / Get screen brightness
# ══════════════════════════════════════════════════════════════════════════════

def tool_set_brightness(action: str, level: Optional[int] = None) -> Dict[str, Any]:
    """
    Controls screen brightness.
    action: 'set' (level 0-100), 'up', 'down', 'get'
    """
    action = action.lower().strip()
    try:
        # Primary: screen-brightness-control package
        try:
            import screen_brightness_control as sbc

            if action == "get":
                current = sbc.get_brightness()
                val = current[0] if isinstance(current, list) else current
                return {"status": "success", "brightness": val, "message": f"Current brightness: {val}%."}

            elif action == "set" and level is not None:
                sbc.set_brightness(max(0, min(100, level)))
                return {"status": "success", "message": f"Brightness set to {level}%."}

            elif action in ("up", "increase"):
                current = sbc.get_brightness()
                cur = current[0] if isinstance(current, list) else current
                new_val = min(100, cur + 10)
                sbc.set_brightness(new_val)
                return {"status": "success", "message": f"Brightness increased to {new_val}%."}

            elif action in ("down", "decrease"):
                current = sbc.get_brightness()
                cur = current[0] if isinstance(current, list) else current
                new_val = max(0, cur - 10)
                sbc.set_brightness(new_val)
                return {"status": "success", "message": f"Brightness decreased to {new_val}%."}

        except ImportError:
            pass

        # Fallback: PowerShell WMI
        if action == "set" and level is not None:
            ps = f"(Get-WmiObject -Namespace root/WMI -Class WmiMonitorBrightnessMethods).WmiSetBrightness(1,{level})"
            result = subprocess.run(["powershell", "-c", ps], capture_output=True, timeout=5)
            if result.returncode == 0:
                return {"status": "success", "message": f"Brightness set to {level}%."}
        elif action == "get":
            ps = "(Get-WmiObject -Namespace root/WMI -Class WmiMonitorBrightness).CurrentBrightness"
            result = subprocess.run(["powershell", "-c", ps], capture_output=True, text=True, timeout=5)
            val = result.stdout.strip()
            if val:
                return {"status": "success", "brightness": int(val), "message": f"Current brightness: {val}%."}

        return {"status": "error", "message": "Could not control brightness. Make sure screen-brightness-control is installed."}

    except Exception as e:
        return {"status": "error", "message": f"Brightness control failed: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
#  6.  WI-FI CONTROL  —  Status, Connect, Disconnect, List Networks
# ══════════════════════════════════════════════════════════════════════════════

def tool_wifi_control(action: str, ssid: str = "") -> Dict[str, Any]:
    """
    Controls Wi-Fi.
    action: 'status', 'list', 'connect' (ssid required), 'disconnect', 'enable', 'disable'
    """
    action = action.lower().strip()
    try:
        if action == "status":
            result = subprocess.run(
                ["netsh", "wlan", "show", "interfaces"],
                capture_output=True, text=True, timeout=8
            )
            output = result.stdout
            if "State" in output:
                for line in output.splitlines():
                    if "State" in line:
                        state = line.split(":")[1].strip()
                    if "SSID" in line and "BSSID" not in line:
                        connected_ssid = line.split(":")[1].strip()
                return {
                    "status": "success",
                    "message": f"Wi-Fi is {state}. Connected to: '{connected_ssid}'."
                }
            return {"status": "success", "message": "Wi-Fi adapter found but not connected to any network."}

        elif action == "list":
            result = subprocess.run(
                ["netsh", "wlan", "show", "availablenetworks"],
                capture_output=True, text=True, timeout=10
            )
            networks = []
            for line in result.stdout.splitlines():
                if "SSID" in line and "BSSID" not in line:
                    net = line.split(":")[1].strip()
                    if net:
                        networks.append(net)
            net_list = ", ".join(networks[:10]) if networks else "No networks found"
            return {"status": "success", "networks": networks, "message": f"Available Wi-Fi networks: {net_list}"}

        elif action == "connect":
            if not ssid:
                return {"status": "error", "message": "Please specify the Wi-Fi network name (SSID) to connect."}
            result = subprocess.run(
                ["netsh", "wlan", "connect", f"name={ssid}"],
                capture_output=True, text=True, timeout=10
            )
            if "successfully" in result.stdout.lower():
                return {"status": "success", "message": f"Connected to Wi-Fi: '{ssid}'."}
            return {"status": "error", "message": f"Could not connect to '{ssid}'. Make sure you've connected before (profile must exist)."}

        elif action == "disconnect":
            subprocess.run(["netsh", "wlan", "disconnect"], capture_output=True, timeout=8)
            return {"status": "success", "message": "Disconnected from Wi-Fi."}

        elif action == "disable":
            subprocess.run(
                ["netsh", "interface", "set", "interface", "Wi-Fi", "disabled"],
                capture_output=True, timeout=8, shell=True
            )
            return {"status": "success", "message": "Wi-Fi disabled."}

        elif action == "enable":
            subprocess.run(
                ["netsh", "interface", "set", "interface", "Wi-Fi", "enabled"],
                capture_output=True, timeout=8, shell=True
            )
            return {"status": "success", "message": "Wi-Fi enabled."}

        return {"status": "error", "message": f"Unknown Wi-Fi action '{action}'. Use: status, list, connect, disconnect, enable, disable"}

    except Exception as e:
        return {"status": "error", "message": f"Wi-Fi control failed: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
#  7.  DOCUMENTS  —  Create Word (.docx) and Excel (.xlsx) files
# ══════════════════════════════════════════════════════════════════════════════

def tool_create_document(
    doc_type: str,
    filename: str,
    content: str = "",
    folder: str = "Documents"
) -> Dict[str, Any]:
    """
    Creates a Word (.docx) or Excel (.xlsx) document and opens it.
    doc_type: 'word' or 'excel'
    filename: name without extension
    content: text for Word, or comma-separated rows for Excel (e.g. 'Name,Age\nAli,25')
    folder: 'Desktop', 'Documents', or full path
    """
    folder_map = {
        "desktop":   Path.home() / "Desktop",
        "documents": Path.home() / "Documents",
        "downloads": Path.home() / "Downloads",
    }
    target_dir = folder_map.get(folder.lower(), Path(folder))
    target_dir.mkdir(parents=True, exist_ok=True)

    safe_name = "".join(c for c in filename if c.isalnum() or c in " _-").strip()
    doc_type = doc_type.lower().strip()

    try:
        if doc_type == "word":
            try:
                from docx import Document
                from docx.shared import Pt
            except ImportError:
                return {"status": "error", "message": "python-docx not installed. Run: pip install python-docx"}

            doc = Document()
            doc.add_heading(safe_name, level=1)

            if content:
                for line in content.splitlines():
                    if line.strip():
                        doc.add_paragraph(line)

            file_path = target_dir / f"{safe_name}.docx"
            doc.save(str(file_path))
            os.startfile(str(file_path))
            log_action("CREATE_WORD", {"file": str(file_path)}, "SUCCESS")
            return {
                "status": "success",
                "path": str(file_path),
                "message": f"Word document '{safe_name}.docx' created and opened in {target_dir.name}."
            }

        elif doc_type in ("excel", "spreadsheet"):
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font
            except ImportError:
                return {"status": "error", "message": "openpyxl not installed. Run: pip install openpyxl"}

            wb = Workbook()
            ws = wb.active
            ws.title = safe_name[:30]

            if content:
                rows = content.strip().splitlines()
                for r_idx, row_text in enumerate(rows, start=1):
                    cells = [c.strip() for c in row_text.split(",")]
                    for c_idx, cell_val in enumerate(cells, start=1):
                        ws.cell(row=r_idx, column=c_idx, value=cell_val)
                    # Bold header row
                    if r_idx == 1:
                        for c_idx in range(1, len(cells) + 1):
                            ws.cell(row=1, column=c_idx).font = Font(bold=True)
            else:
                # Empty spreadsheet with default headers
                ws.cell(row=1, column=1, value="Column A").font = Font(bold=True)
                ws.cell(row=1, column=2, value="Column B").font = Font(bold=True)

            file_path = target_dir / f"{safe_name}.xlsx"
            wb.save(str(file_path))
            os.startfile(str(file_path))
            log_action("CREATE_EXCEL", {"file": str(file_path)}, "SUCCESS")
            return {
                "status": "success",
                "path": str(file_path),
                "message": f"Excel spreadsheet '{safe_name}.xlsx' created and opened in {target_dir.name}."
            }

        else:
            return {"status": "error", "message": f"Unknown document type '{doc_type}'. Use: word or excel"}

    except Exception as e:
        return {"status": "error", "message": f"Failed to create {doc_type} document: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
#  8.  NEWS  —  Latest headlines from Google News RSS (no API key)
# ══════════════════════════════════════════════════════════════════════════════

def tool_get_news(topic: str = "", count: int = 5, language: str = "en") -> Dict[str, Any]:
    """
    Fetches latest news headlines via Google News RSS — no API key needed.
    topic: empty for top stories, or 'technology', 'sports', 'business', etc.
    """
    try:
        import feedparser

        if topic.strip():
            query = urllib.parse.quote(topic.strip())
            rss_url = f"https://news.google.com/rss/search?q={query}&hl={language}&gl=IN&ceid=IN:{language.upper()}"
        else:
            rss_url = f"https://news.google.com/rss?hl={language}&gl=IN&ceid=IN:{language.upper()}"

        feed = feedparser.parse(rss_url)
        entries = feed.entries[:count]

        if not entries:
            return {"status": "success", "headlines": [], "message": "No news found. Check your internet connection."}

        headlines = []
        for entry in entries:
            title   = entry.get("title", "No Title").split(" - ")[0].strip()  # Remove source suffix
            source  = entry.get("source", {}).get("title", "Unknown") if hasattr(entry, "source") else ""
            pub_date = entry.get("published", "")[:16]
            headlines.append({"title": title, "source": source, "date": pub_date})

        news_lines = [f"{i+1}. {h['title']}" for i, h in enumerate(headlines)]
        summary    = f"Top {len(headlines)} {'news about ' + topic if topic else 'headlines'}: " + "; ".join(news_lines)

        log_action("GET_NEWS", {"topic": topic, "count": len(headlines)}, "SUCCESS")
        return {
            "status": "success",
            "headlines": headlines,
            "message": summary
        }
    except ImportError:
        return {"status": "error", "message": "feedparser not installed. Run: pip install feedparser"}
    except Exception as e:
        return {"status": "error", "message": f"Failed to fetch news: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
#  9.  MORNING BRIEFING  —  Weather + News + Battery + Time
# ══════════════════════════════════════════════════════════════════════════════

def tool_morning_briefing(city: str = "") -> Dict[str, Any]:
    """
    Full morning briefing: time, battery status, today's weather, and top 3 news headlines.
    Speaks everything aloud automatically.
    """
    parts = []
    now = datetime.datetime.now()
    greeting = "Good morning" if now.hour < 12 else ("Good afternoon" if now.hour < 17 else "Good evening")
    parts.append(f"{greeting}! It's {now.strftime('%A, %d %B %Y')} and the time is {now.strftime('%I:%M %p')}.")

    # Battery
    try:
        import psutil
        bat = psutil.sensors_battery()
        if bat:
            charge_str = "charging" if bat.power_plugged else "not charging"
            parts.append(f"Your battery is at {round(bat.percent)}%, {charge_str}.")
    except Exception:
        pass

    # Weather
    try:
        from core.system_tools import tool_get_weather
        weather = tool_get_weather(city)
        if weather.get("status") == "success":
            parts.append(weather["message"])
    except Exception:
        pass

    # News
    try:
        news = tool_get_news(count=3)
        if news.get("status") == "success" and news.get("headlines"):
            headlines = [h["title"] for h in news["headlines"]]
            parts.append("Today's top news: " + ". ".join(headlines) + ".")
    except Exception:
        pass

    full_briefing = " ".join(parts)

    # Speak it aloud
    def _speak_briefing(text: str) -> None:
        try:
            from voice.tts import TextToSpeechEngine
            TextToSpeechEngine().speak(text)
        except Exception:
            pass

    threading.Thread(target=_speak_briefing, args=(full_briefing,), daemon=True).start()
    log_action("MORNING_BRIEFING", {"city": city}, "SUCCESS")

    return {
        "status": "success",
        "briefing": full_briefing,
        "message": full_briefing
    }


# ══════════════════════════════════════════════════════════════════════════════
#  10. SMART HOME  —  TP-Link Kasa + Philips Hue
# ══════════════════════════════════════════════════════════════════════════════

def tool_smart_home(device: str, action: str, value: Optional[int] = None) -> Dict[str, Any]:
    """
    Controls smart home devices.
    device: 'light', 'ac', 'fan', 'plug', or specific device name
    action: 'on', 'off', 'dim' (value 0-100), 'toggle', 'status'
    Supports: TP-Link Kasa and Philips Hue
    """
    device_lower = device.lower().strip()
    action_lower = action.lower().strip()

    # ── TP-Link Kasa ─────────────────────────────────────────────────────────
    if KASA_DEVICE_IP:
        try:
            import asyncio

            async def _kasa_control():
                try:
                    from kasa import SmartBulb, SmartPlug
                    # Try as bulb first, then plug
                    try:
                        dev = SmartBulb(KASA_DEVICE_IP)
                    except Exception:
                        dev = SmartPlug(KASA_DEVICE_IP)

                    await dev.update()
                    if action_lower == "on":
                        await dev.turn_on()
                    elif action_lower == "off":
                        await dev.turn_off()
                    elif action_lower == "toggle":
                        await dev.turn_off() if dev.is_on else await dev.turn_on()
                    elif action_lower == "dim" and value is not None:
                        if hasattr(dev, "set_brightness"):
                            await dev.set_brightness(max(0, min(100, value)))
                    await dev.update()
                    state = "on" if dev.is_on else "off"
                    return {"status": "success", "message": f"Kasa {device}: turned {state}."}
                except ImportError:
                    return None
                except Exception as e:
                    return {"status": "error", "message": f"Kasa error: {str(e)}"}

            result = asyncio.run(_kasa_control())
            if result:
                log_action("SMART_HOME_KASA", {"device": device, "action": action}, "SUCCESS")
                return result
        except Exception:
            pass

    # ── Philips Hue ──────────────────────────────────────────────────────────
    if HUE_BRIDGE_IP and HUE_USERNAME:
        try:
            import urllib.request
            # Get all lights
            api_base = f"http://{HUE_BRIDGE_IP}/api/{HUE_USERNAME}"
            lights_url = f"{api_base}/lights"
            req = urllib.request.urlopen(lights_url, timeout=5)
            lights = json.loads(req.read())

            if not lights:
                return {"status": "error", "message": "No Hue lights found on bridge."}

            # Find matching light (first one if device is generic)
            target_id = None
            for lid, ldata in lights.items():
                lname = ldata.get("name", "").lower()
                if device_lower in lname or device_lower in ("light", "lights", "all"):
                    target_id = lid
                    break

            if not target_id:
                target_id = list(lights.keys())[0]  # fallback: first light

            state_url = f"{api_base}/lights/{target_id}/state"
            payload = {}
            if action_lower == "on":
                payload = {"on": True}
            elif action_lower == "off":
                payload = {"on": False}
            elif action_lower == "dim" and value is not None:
                payload = {"on": True, "bri": int(254 * value / 100)}
            elif action_lower == "toggle":
                current = lights[target_id].get("state", {}).get("on", False)
                payload = {"on": not current}

            if payload:
                data = json.dumps(payload).encode("utf-8")
                req = urllib.request.Request(state_url, data=data, method="PUT")
                urllib.request.urlopen(req, timeout=5)
                log_action("SMART_HOME_HUE", {"device": device, "action": action}, "SUCCESS")
                return {"status": "success", "message": f"Hue {device}: {action} executed."}

        except Exception as e:
            return {"status": "error", "message": f"Philips Hue error: {str(e)}"}

    # ── Not configured ────────────────────────────────────────────────────────
    return {
        "status": "not_configured",
        "message": (
            "Smart home devices are not configured yet. "
            "For TP-Link Kasa: set KASA_DEVICE_IP in .env and run 'pip install python-kasa'. "
            "For Philips Hue: set HUE_BRIDGE_IP and HUE_USERNAME in .env."
        )
    }
